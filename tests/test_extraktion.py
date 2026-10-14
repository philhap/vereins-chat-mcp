"""Tests für die Textextraktion. Fixtures werden im Test erzeugt (docx real, PDF leer)."""

from io import BytesIO

from webling_mcp.rag.extraktion import extrahiere_text


def docx_bytes(*absaetze: str) -> bytes:
    from docx import Document

    d = Document()
    for a in absaetze:
        d.add_paragraph(a)
    puffer = BytesIO()
    d.save(puffer)
    return puffer.getvalue()


def test_docx_wird_extrahiert():
    daten = docx_bytes("Protokoll der Sitzung.", "TOP 1: Beitragserhöhung.")
    text = extrahiere_text("protokoll.docx", daten)
    assert "Beitragserhöhung" in text


def test_docx_tabellen_werden_extrahiert():
    from docx import Document

    d = Document()
    tabelle = d.add_table(rows=1, cols=2)
    tabelle.rows[0].cells[0].text = "Posten"
    tabelle.rows[0].cells[1].text = "Betrag"
    puffer = BytesIO()
    d.save(puffer)

    text = extrahiere_text("kasse.docx", puffer.getvalue())
    assert "Posten" in text and "Betrag" in text


def test_pdf_ohne_textebene_liefert_none():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=595, height=842)
    puffer = BytesIO()
    w.write(puffer)

    assert extrahiere_text("scan.pdf", puffer.getvalue()) is None


def test_unbekannte_endung_liefert_none():
    assert extrahiere_text("foto.jpg", b"\xff\xd8\xff") is None


def test_kaputte_datei_liefert_none():
    assert extrahiere_text("kaputt.pdf", b"kein pdf") is None
