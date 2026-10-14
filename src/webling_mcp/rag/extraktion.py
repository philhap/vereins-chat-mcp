"""
Textextraktion für den RAG-Ingest: PDF (pypdf) und Word (python-docx).

Rückgabe None = Dokument kann nicht indexiert werden (unbekannter Typ,
Scan ohne Textebene, defekte Datei) — der Ingest meldet es als übersprungen.
"""

import logging
from io import BytesIO
from pathlib import Path

logger = logging.getLogger(__name__)


def extrahiere_text(dateiname: str, daten: bytes) -> str | None:
    """Extrahiert Klartext aus PDF/DOCX-Bytes; None wenn nicht möglich."""
    endung = Path(dateiname).suffix.lower()
    try:
        if endung == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(daten))
            text = "\n".join((seite.extract_text() or "") for seite in reader.pages)
        elif endung == ".docx":
            from docx import Document

            doc = Document(BytesIO(daten))
            teile = [absatz.text for absatz in doc.paragraphs]
            for tabelle in doc.tables:
                for zeile in tabelle.rows:
                    teile.append(" | ".join(zelle.text for zelle in zeile.cells))
            text = "\n".join(teile)
        else:
            logger.info("Kein Extraktor für %s (%s)", endung, dateiname)
            return None
    except Exception as e:
        logger.warning("Extraktion fehlgeschlagen (%s): %s", dateiname, type(e).__name__)
        return None
    text = text.strip()
    return text or None
